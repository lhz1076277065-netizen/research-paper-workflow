#!/usr/bin/env python3
"""Actual editable exports and adversarial result/render checks.

Requires python-docx, ReportLab and either pdftotext or pypdf. The PDF renderer
reads DOCX paragraphs; LaTeX source text is checked separately. This development
fixture does not test arbitrary journal layouts or replace visual inspection.
"""
from __future__ import annotations
import argparse
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import shutil
import sys
import textwrap

def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def save(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
def ref(path,root):return {'path':str(path.relative_to(root)),'sha256':sha(path)}

def render_word(path,target):
    from docx import Document
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import landscape,A4
    c=canvas.Canvas(str(target),pagesize=landscape(A4));c.setFont('Helvetica',11)
    y=landscape(A4)[1]-45
    for p in Document(path).paragraphs:
        for line in textwrap.wrap(p.text,115):c.drawString(35,y,line);y-=17
    c.save()

def main():
    from docx import Document
    import docx,reportlab
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--digital',required=True);parser.add_argument('--out',required=True)
    args=parser.parse_args();digital=Path(args.digital).resolve();out=Path(args.out).resolve();out.mkdir(parents=True,exist_ok=False)
    helper=Path(__file__).resolve().parents[2]/'src/common/scripts/result_links.py'
    spec=importlib.util.spec_from_file_location('acceptance_links',helper);links=importlib.util.module_from_spec(spec);spec.loader.exec_module(links)
    shutil.copyfile(__file__,out/'executed-export-source.py')
    for name in ['public-results.json','wine-quality.zip','executed-source.py','research-model.npz']:
        shutil.copyfile(digital/name,out/name)
    actual=json.loads((out/'public-results.json').read_text())
    sentence=(f'Ridge regression RMSE = {actual["test_rmse"]:.3f} quality score points on heldout white wine data '
              f'(n={actual["split"]["test"]}), lower than the training mean baseline '
              f'({actual["baseline_test_rmse"]:.3f} quality score points); outcome: wine quality. Measurement time: not supplied.')
    semantic={'unit':'quality score points','denominator':actual['split']['test'],
              'population':'white wine','outcome':'wine quality','comparison':'training mean baseline',
              'model':'ridge regression','split':'heldout','direction':'decrease','time':'not supplied'}
    records={'RMSE-ridge':{'value':actual['test_rmse'],**semantic},
             'RMSE-baseline':{'value':actual['baseline_test_rmse'],'unit':'quality score points','denominator':actual['split']['test']}}
    registry=out/'results.json';save(registry,records)
    md=out/'manuscript.md';md.write_text(sentence+'\n')
    tex=out/'manuscript.tex';tex.write_text('\\documentclass{article}\n\\begin{document}\n'+sentence+'\n\\end{document}\n')
    word=out/'manuscript.docx';d=Document();d.add_paragraph(sentence);d.save(word)
    pdf=out/'manuscript.pdf';render_word(word,pdf)
    sources=[{'id':'actual',**ref(registry,out),'versions':{'data':ref(out/'wine-quality.zip',out),
               'code':ref(out/'executed-source.py',out),'execution':ref(out/'public-results.json',out)}}]
    target_semantics={k:{'value':v,'text':('lower' if k=='direction' else f'n={v}' if k=='denominator' else v)} for k,v in semantic.items()}
    targets=[(md,{'line':1},'body'),(tex,{'line':3},'latex'),(word,{'paragraph':1},'word'),(pdf,{'page':1,'line':1,'line_end':3},'pdf')]
    checks=[]
    for path,locator,role in targets:
        checks.append({'result_id':'RMSE-ridge','source_id':'actual','role':role,'artifact':ref(path,out),
           'locator':locator,'numeric':[{'field':'value','text':f'{actual["test_rmse"]:.3f}','decimals':3}],
           'semantics':copy.deepcopy(target_semantics)})
        checks.append({'result_id':'RMSE-baseline','source_id':'actual','role':role,'artifact':ref(path,out),
           'locator':locator,'numeric':[{'field':'value','text':f'{actual["baseline_test_rmse"]:.3f}','decimals':3}],
           'semantics':{'unit':{'value':'quality score points','text':'quality score points'},
                        'denominator':{'value':actual['split']['test'],'text':f'n={actual["split"]["test"]}'}}})
    payload={'kind':'result-links','sources':sources,'links':checks}
    correct=links.audit_links(payload,out);save(out/'links-correct.json',correct)
    assert correct['passed'],correct
    # Update the edited artifact fingerprint so each failure is a content error,
    # rather than merely a file hash mismatch.
    mutations={'value':(f'{actual["test_rmse"]:.3f}','0.650'),'unit':('quality score points','seconds'),
               'direction':('lower','higher'),'denominator':(f'n={actual["split"]["test"]}','n=914'),
               'time':('not supplied','6 months')}
    negative={}
    for kind,(old,new) in mutations.items():
        bad=out/f'wrong-{kind}.md';bad.write_text(sentence.replace(old,new)+'\n')
        changed=copy.deepcopy(payload);changed['links']=copy.deepcopy(checks[:1]);changed['links'][0]['artifact']=ref(bad,out)
        report=links.audit_links(changed,out);assert not report['passed'] and report['errors'],report
        negative[kind]=report
    save(out/'intentional-mismatches.json',negative)
    manifest={'kind':'render-dependencies','artifacts':[{'id':key,**ref(path,out)} for key,path in [('md',md),('word',word),('tex',tex),('pdf',pdf)]],
              'renders':[{'output':'word','inputs':{'md':sha(md)},'renderer':{'name':'python-docx','version':docx.__version__}},
                         {'output':'tex','inputs':{'md':sha(md)},'renderer':{'name':'plain LaTeX source exporter','version':'1'}},
                         {'output':'pdf','inputs':{'word':sha(word)},'renderer':{'name':'DOCX paragraph → ReportLab PDF','version':reportlab.Version}}]}
    assert links.audit_render(manifest,out)['passed']
    md.write_text(sentence+'\nAuthor identity and declaration are pending.\n')
    stale=links.audit_render(manifest,out);save(out/'stale-exports.json',stale)
    assert not stale['passed'] and {'word','tex','pdf'}<=set(stale['rebuild_order']),stale
    d=Document();d.add_paragraph(sentence);d.add_paragraph('Author identity and declaration are pending.');d.save(word)
    tex.write_text('\\documentclass{article}\n\\begin{document}\n'+md.read_text()+'\\end{document}\n')
    render_word(word,pdf)
    for item in manifest['artifacts']:item['sha256']=sha(out/item['path'])
    for step in manifest['renders']:
        for key in step['inputs']:step['inputs'][key]=next(x['sha256'] for x in manifest['artifacts'] if x['id']==key)
    clean=links.audit_render(manifest,out);assert clean['passed'],clean
    for check in payload['links']:check['artifact']['sha256']=sha(out/check['artifact']['path'])
    final=links.audit_links(payload,out);assert final['passed'],final
    save(out/'result-links.json',payload);save(out/'render-manifest.json',manifest)
    save(out/'links-final.json',final);save(out/'render-final.json',clean)
    report={'actual_formats':['Markdown','LaTeX source text','DOCX paragraphs','PDF extracted text'],
            'linked_occurrences':len(checks),'intentional_mismatches_detected':list(negative),'rebuild_order':stale['rebuild_order'],
            'rebuilt_outputs':['manuscript.docx','manuscript.tex','manuscript.pdf'],'final_links_passed':final['passed'],
            'final_render_passed':clean['passed'],'latex_compilation_evaluated_here':False,'visual_inspection_performed_here':False,
            'renderer_scope':'Actual DOCX paragraph text rendering with ReportLab; not arbitrary Word/LaTeX layout preservation',
            'code_sha256':sha(__file__)}
    save(out/'acceptance.json',report);print(json.dumps(report,indent=2))

if __name__=='__main__':main()
