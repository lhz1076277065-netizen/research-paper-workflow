"""Create an editable Word report from the canonical polished manuscript."""
import pathlib,re,json,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=pathlib.Path(__file__).resolve().parent;out=R/'delivery';out.mkdir(exist_ok=True)
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.68);sec.left_margin=sec.right_margin=Inches(.75)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
for name,size in [('Title',21),('Heading 1',14),('Heading 2',11)]:
    s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(7)
doc.styles['Title'].paragraph_format.space_before=Pt(0)
doc.styles['Caption'].font.name='Arial';doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.color.rgb=RGBColor(0,0,0)
doc.styles['Caption'].font.bold=False
# The bundled default template has a decorative Title border. Remove it so
# typography and spacing, rather than a horizontal rule, establish hierarchy.
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.core_properties.title='Training only standardization and fixed distance classifiers in Wine and Iris'
doc.core_properties.subject='Reproducible empirical teaching reanalysis'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
lines=(R/'manuscript-final.md').read_text().splitlines();i=0;table_count=0;images=0
while i<len(lines):
    line=lines[i]
    if not line.strip():i+=1;continue
    if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[x.strip() for x in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',v) for v in cells):rows.append(cells)
            i+=1
        table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1';table.autofit=False
        widths=[.55,2.00,1.5,1.05,1.55]
        for j,v in enumerate(rows[0]):table.rows[0].cells[j].text=v
        prop=table.rows[0]._tr.get_or_add_trPr();prop.append(OxmlElement('w:tblHeader'))
        for vals in rows[1:]:
            cells=table.add_row().cells
            for j,v in enumerate(vals):cells[j].text=v
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    for run in p.runs:
                        run.font.size=Pt(8.5);run.font.color.rgb=RGBColor(0,0,0)
        table_count+=1;continue
    elif line.startswith('!['):
        alt,path=re.match(r'!\[(.*?)\]\((.*?)\)',line).groups()
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
        pic=p.add_run().add_picture(str(R/path),width=Inches(6))
        pic._inline.docPr.set('descr',alt+'; full descriptive caption follows')
        images+=1
    else:
        p=doc.add_paragraph(line,'Caption' if line.startswith(('Figure 1.','Figure 2.','Table 1.')) else 'Normal')
        p.paragraph_format.widow_control=True
    i+=1
dest=out/'Wine_Iris_teaching_reanalysis.docx';doc.save(dest)
with zipfile.ZipFile(dest) as z:
    assert z.testzip() is None
    xml=z.read('word/document.xml').decode()
    assert 'Training only standardization' in xml and '25.82' in xml and '-7.00' in xml
    assert xml.count('<w:tbl>')==1 and xml.count('<wp:docPr ')==2
report={'canonical_markdown':'manuscript-final.md','docx':str(dest.relative_to(R)),'table_count':table_count,'embedded_figures':images,'archive_crc_passed':True,'author_not_invented':True,'render_required':True}
(R/'document-structure-check.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
