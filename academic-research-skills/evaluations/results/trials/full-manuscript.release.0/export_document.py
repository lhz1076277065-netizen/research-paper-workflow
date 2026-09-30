"""Export the complete Markdown manuscript as an editable DOCX."""
from pathlib import Path
import hashlib
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
source = OUT / "manuscript.md"
text = source.read_text()
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(1)
normal = document.styles["Normal"]
normal.font.name, normal.font.size = "Times New Roman", Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for name, size in [("Title", 16), ("Heading 1", 12), ("Heading 2", 11)]:
    style = document.styles[name]
    style.font.name, style.font.size = "Arial", Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(25, 25, 25)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)
header = section.header.paragraphs[0]
header.add_run("Fixed synthetic development sample").font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("Page ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

def plain(value):
    return re.sub(r"\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`", lambda m: next(g for g in m.groups() if g is not None), value)

def add_inline(paragraph, value):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", value):
        run = paragraph.add_run(plain(part))
        if part.startswith("**"):
            run.bold = True
        elif part.startswith("*"):
            run.italic = True

expected_units = []
lines = text.splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1
        continue
    if line.startswith("|"):
        raw_rows = []
        while i < len(lines) and lines[i].startswith("|"):
            raw_rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        rows = [row for row in raw_rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
        table = document.add_table(rows=0, cols=len(rows[0]))
        table.style = "Light Shading Accent 1"
        table.autofit = False
        widths = [1.03, 0.55, 0.60, 0.45, 0.55, 3.12]
        for row_index, values in enumerate(rows):
            row = table.add_row()
            props = row._tr.get_or_add_trPr()
            no_split = OxmlElement("w:cantSplit")
            props.append(no_split)
            if row_index == 0:
                repeated = OxmlElement("w:tblHeader")
                props.append(repeated)
            for col_index, value in enumerate(values):
                cell = row.cells[col_index]
                cell.width = Inches(widths[col_index])
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                add_inline(paragraph, value)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = row_index == 0
                expected_units.append(plain(value))
        continue
    image = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
    if image:
        document.add_page_break()
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        picture = paragraph.add_run().add_picture(str(OUT / image.group(2)), width=Inches(6.3), height=Inches(4.7))
        picture._inline.docPr.set("descr", "Point-and-interval display of all eight manufactured summary losses and four weighted-average points. B is lower in S1, higher in S2; its reference-average advantage reverses under the changed condition. Supplied intervals belong only to summary rows.")
        i += 1
        continue
    heading = re.match(r"^(#{1,3}) (.*)$", line)
    if heading:
        depth = len(heading.group(1))
        value = heading.group(2)
        paragraph = document.add_paragraph(style="Title" if depth == 1 else f"Heading {depth-1}")
        add_inline(paragraph, value)
        expected_units.append(plain(value))
        i += 1
        continue
    if line.startswith("- "):
        value = line[2:]
        paragraph = document.add_paragraph(style="List Bullet")
        add_inline(paragraph, value)
        expected_units.append(plain(value))
        i += 1
        continue
    paragraph = document.add_paragraph()
    add_inline(paragraph, line)
    if line.startswith("**Table 1."):
        paragraph.paragraph_format.keep_with_next = True
    if line.startswith("**Figure 1."):
        for run in paragraph.runs:
            run.font.size = Pt(9.5)
    expected_units.append(plain(line))
    i += 1

digest = hashlib.sha256(source.read_bytes()).hexdigest()
document.core_properties.author = ""
document.core_properties.last_modified_by = "Codex"
document.core_properties.title = plain(lines[0][2:])
document.core_properties.subject = "Fixed synthetic manuscript revision"
document.core_properties.comments = f"Complete editable manuscript derived from manuscript.md SHA-256 {digest}; no empirical study or submission."
document.save(OUT / "manuscript.docx")
(OUT / "process/document-export.json").write_text(json.dumps({"source_sha256": digest, "source": "manuscript.md", "expected_text_units": expected_units, "figure_size_inches": [6.3, 4.7], "exporter": "python-docx 1.2.0"}, indent=2) + "\n")
print(json.dumps({"docx": str(OUT / "manuscript.docx"), "text_units": len(expected_units), "source_sha256": digest}))
