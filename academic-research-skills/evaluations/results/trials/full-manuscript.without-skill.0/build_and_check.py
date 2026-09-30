"""Reproduce displays, export Word, and check the fixed synthetic revision.

Run with the shared fixture Python for --figure; bundled document Python for
--docx and --check. No new evidence or statistical inference is generated.
"""
from pathlib import Path
from decimal import Decimal
import argparse
import csv
import hashlib
import io
import json
import re
import sys

OUT = Path(__file__).resolve().parent
REQUEST = Path('LOCAL_EVIDENCE_ROOT/evaluation-suite/requests/full-manuscript.without-skill.0.json')


def load():
    request = json.loads(REQUEST.read_text())
    receipts = []
    for item in request['material_files']:
        actual = hashlib.sha256(Path(item['path']).read_bytes()).hexdigest()
        assert actual == item['sha256'], item['path']
        receipts.append({**item, 'actual_sha256': actual, 'matches': True})
    csv_path = Path(next(x['path'] for x in request['material_files'] if x['path'].endswith('.csv')))
    rows = list(csv.DictReader(io.StringIO('\n'.join(x for x in csv_path.read_text().splitlines() if not x.startswith('#')))))
    assert len(rows) == 8
    averages = {}
    denominators = {}
    for condition in ('reference', 'changed'):
        averages[condition] = {}
        denominators[condition] = {}
        for method in ('A', 'B'):
            subset = [r for r in rows if r['condition'] == condition and r['method'] == method]
            denominator = sum(int(r['n']) for r in subset)
            assert denominator == 100
            denominators[condition][method] = denominator
            averages[condition][method] = sum(Decimal(r['n']) * Decimal(r['loss']) for r in subset) / denominator
    assert averages == {'reference': {'A': Decimal('4.80'), 'B': Decimal('4.36')}, 'changed': {'A': Decimal('6.60'), 'B': Decimal('7.52')}}
    assert averages['reference']['B'] - averages['reference']['A'] == Decimal('-0.44')
    assert averages['changed']['B'] - averages['changed']['A'] == Decimal('0.92')
    return rows, averages, receipts, denominators


def figure(rows, averages):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.size': 10, 'font.family': 'DejaVu Sans', 'svg.hashsalt': 'fixed-manuscript-fixture', 'svg.fonttype': 'none', 'pdf.fonttype': 42})
    fig, axes = plt.subplots(3, 1, figsize=(7.4, 7.0), sharex=True)
    colors = {'A': '#17608A', 'B': '#B54D17'}
    markers = {'A': 'o', 'B': 's'}
    offsets = {'A': .13, 'B': -.13}
    for index, condition in enumerate(('reference', 'changed')):
        ax = axes[index]
        subset = [r for r in rows if r['condition'] == condition]
        for r in subset:
            method = r['method']
            y = (1 if r['stratum'] == 'S1' else 0) + offsets[method]
            loss, lo, hi = (float(r[x]) for x in ('loss', 'lo', 'hi'))
            ax.errorbar(loss, y, xerr=[[loss - lo], [hi - loss]], fmt=markers[method], color=colors[method], markersize=6, capsize=4, lw=1.5, label=method)
            ax.text(loss, y + (.12 if method == 'A' else -.18), f'{loss:.1f}', ha='center', va='center', fontsize=9, color=colors[method])
        counts = {r['stratum']: r['n'] for r in subset}
        ax.set_yticks([1, 0], [f"S1 (n = {counts['S1']})", f"S2 (n = {counts['S2']})"])
        ax.set_title(f"{'AB'[index]}  {condition.capitalize()} strata — supplied row 95% intervals", loc='left', fontweight='bold', fontsize=10)
    ax = axes[2]
    for condition, y0 in [('reference', 1), ('changed', 0)]:
        for method in ('A', 'B'):
            value = float(averages[condition][method])
            y = y0 + offsets[method]
            ax.plot(value, y, markers[method], color=colors[method], markersize=6)
            ax.text(value, y + (.12 if method == 'A' else -.18), f'{value:.2f}', ha='center', va='center', fontsize=9, color=colors[method])
        delta = averages[condition]['B'] - averages[condition]['A']
        ax.text(9.75, y0, f'B − A = {delta:+.2f}', ha='right', va='center', fontsize=9)
    ax.set_yticks([1, 0], ['Reference', 'Changed'])
    ax.set_title('C  Weighted averages — no aggregate intervals', loc='left', fontweight='bold', fontsize=10)
    ax.set_xlabel('Loss (lower is better; units unspecified)')
    for ax in axes:
        ax.set_xlim(0, 10.5)
        ax.set_ylim(-.45, 1.48)
        ax.set_xticks(range(0, 11, 2))
        ax.grid(axis='x', color='#DDDDDD', lw=.6)
        ax.set_axisbelow(True)
        for side in ('top', 'right', 'left'):
            ax.spines[side].set_visible(False)
        ax.tick_params(axis='y', length=0)
    from matplotlib.lines import Line2D
    fig.legend([Line2D([0], [0], marker=markers[m], color=colors[m], linestyle='none', markersize=6) for m in ('A', 'B')], ['Method A', 'Method B'], loc='upper center', ncol=2, frameon=False, bbox_to_anchor=(.59, .995))
    fig.subplots_adjust(top=.92, bottom=.09, left=.20, right=.98, hspace=.52)
    for suffix in ('png', 'svg', 'pdf'):
        fig.savefig(OUT / f'figure1.{suffix}', dpi=300, metadata={'Creator': 'Fixed synthetic manuscript revision'})
    plt.close(fig)
    (OUT / 'display-data.json').write_text(json.dumps({'display': 'Figure 1', 'row_intervals': rows, 'weighted_averages': {c: {m: str(v) for m, v in vals.items()} for c, vals in averages.items()}, 'aggregate_intervals': None, 'transformation': 'sum(n*loss)/sum(n), independently by condition and method; supplied row intervals unchanged'}, indent=2) + '\n')


def plain(text):
    return text.replace('*', '').replace('`', '').strip()


def blocks(markdown):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            table = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [plain(x.strip()) for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                    table.append(cells)
                i += 1
            yield 'table', table
            continue
        if line.startswith('!['):
            alt, path = re.fullmatch(r'!\[(.*)\]\((.*)\)', line).groups()
            yield 'image', (alt, path)
        elif line.startswith('#'):
            level, title = line.split(' ', 1)
            yield 'heading', (len(level), title)
        else:
            paragraph = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not lines[i + 1].strip().startswith(('#', '|', '![', '- ')):
                i += 1
                paragraph.append(lines[i].strip())
            yield 'paragraph', plain(' '.join(paragraph))
        i += 1


def docx():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.9)
    normal = document.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.keep_together = True
    for style, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11)]:
        document.styles[style].font.name = 'Aptos'
        document.styles[style].font.size = Pt(size)
        document.styles[style].font.color.rgb = RGBColor.from_string('183C50')
        document.styles[style].paragraph_format.keep_with_next = True
    section.header.paragraphs[0].text = 'Fixed synthetic development sample'
    section.header.paragraphs[0].style = document.styles['Caption']
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for kind, content in blocks((OUT / 'manuscript.md').read_text()):
        if kind == 'heading':
            level, title = content
            paragraph = document.add_paragraph(title, 'Title' if level == 1 else f'Heading {level - 1}')
            if title.startswith('Figure 1.') or title == 'Discussion':
                paragraph.paragraph_format.page_break_before = True
        elif kind == 'table':
            table = document.add_table(rows=0, cols=len(content[0]))
            table.style = 'Light Shading Accent 1'
            for index, row in enumerate(content):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(4)
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            run.bold = index == 0
                if index == 0:
                    repeat = OxmlElement('w:tblHeader')
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
        elif kind == 'image':
            alt, path = content
            picture = document.add_picture(str(OUT / path), width=Inches(6.35))
            picture._inline.docPr.set('descr', alt)
        else:
            paragraph = document.add_paragraph(content)
            if content.startswith('Caption.'):
                paragraph.style = document.styles['Caption']
                paragraph.paragraph_format.keep_with_next = False
    document.core_properties.title = 'Conditional average gains in estimation'
    document.core_properties.subject = 'Fixed synthetic manuscript development sample'
    document.core_properties.author = ''
    document.save(OUT / 'manuscript.docx')


def check(rows, averages, receipts, denominators):
    from docx import Document
    markdown = (OUT / 'manuscript.md').read_text()
    parsed = list(blocks(markdown))
    expected_rows = [[r['condition'].capitalize(), r['stratum'], r['method'], r['n'], r['loss'], f"[{r['lo']}, {r['hi']}]"] for r in rows]
    tables = [x for kind, x in parsed if kind == 'table']
    assert len(tables) == 1 and tables[0][1:] == expected_rows
    document = Document(OUT / 'manuscript.docx')
    expected_paragraphs = [content[1] if kind == 'heading' else content for kind, content in parsed if kind in ('heading', 'paragraph')]
    actual_paragraphs = [p.text for p in document.paragraphs if p.text]
    assert actual_paragraphs == expected_paragraphs, 'Markdown and Word text differ'
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == tables[0]
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0]._inline.docPr.get('descr')
    display = json.loads((OUT / 'display-data.json').read_text())
    assert display['row_intervals'] == rows and display['aggregate_intervals'] is None
    assert display['weighted_averages'] == {c: {m: str(v) for m, v in vals.items()} for c, vals in averages.items()}
    headings = [content[1] for kind, content in parsed if kind == 'heading']
    for section in ('Abstract', 'Introduction', 'Methods', 'Results', 'Discussion', 'Conclusion', 'Declarations and references'):
        assert section in headings
    for filename in ('figure1.png', 'figure1.svg', 'figure1.pdf', 'factual_claim_comparison.md', 'operation_note.md'):
        assert (OUT / filename).is_file() and (OUT / filename).stat().st_size > 0
    report = {'status': 'passed', 'review_type': 'local agent checks, not human scientific approval', 'material_sha256': receipts, 'row_count': len(rows), 'row_values_counts_intervals_preserved': True, 'weighted_losses': {c: {m: str(v) for m, v in vals.items()} for c, vals in averages.items()}, 'differences_B_minus_A': {c: str(v['B'] - v['A']) for c, v in averages.items()}, 'condition_method_denominators': denominators, 'all_manuscript_sections_present': True, 'docx_markdown_text_and_table_parity': True, 'figure_inputs_match_csv': True, 'aggregate_intervals_created': False, 'image_alt_text_present': True, 'script_runtime': sys.version}
    (OUT / 'verification.json').write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--figure', action='store_true')
    parser.add_argument('--docx', action='store_true')
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    if not (args.figure or args.docx or args.check):
        parser.error('Choose --figure, --docx or --check')
    rows, averages, receipts, denominators = load()
    if args.figure:
        figure(rows, averages)
    if args.docx:
        docx()
    if args.check:
        check(rows, averages, receipts, denominators)
