import csv,json,os,runpy,shutil,subprocess,sys,time
from pathlib import Path
out=Path(__file__).parent.resolve();receipts=[]
def source(uid):return Path(json.loads((out/'prepared'/f'{uid}.json').read_text())['root'])
def run(uid,rel,args,cwd=None):
 cmd=[sys.executable,str(source(uid)/rel),*map(str,args)];s=time.monotonic();p=subprocess.run(cmd,capture_output=True,text=True,timeout=30,cwd=cwd or out,env=dict(os.environ,MPLBACKEND='Agg'))
 (out/f'{uid}-actual-stdout.log').write_text(p.stdout);(out/f'{uid}-actual-stderr.log').write_text(p.stderr)
 receipts.append(dict(id=uid,command=cmd,returncode=p.returncode,seconds=time.monotonic()-s));return p
# Original, unmodified Rougier code with its relative output directory preserved in working dir.
p=out/'rougier';(p/'code/ornaments').mkdir(parents=True,exist_ok=True);(p/'figures/ornaments').mkdir(parents=True,exist_ok=True)
cmd=[sys.executable,'-c',f'import runpy,matplotlib.pyplot as plt;runpy.run_path({str(source("rougier-reference")/"code/ornaments/legend-alternatives.py")!r});plt.gcf().savefig({str(p/"legend-preview.png")!r},dpi=150)']
r=subprocess.run(cmd,cwd=p/'code/ornaments',capture_output=True,text=True,env=dict(os.environ,MPLBACKEND='Agg'),timeout=30)
(out/'rougier-stdout.log').write_text(r.stdout);(out/'rougier-stderr.log').write_text(r.stderr);assert r.returncode==0
receipts.append(dict(id='rougier-reference',command=cmd,returncode=r.returncode))
# Image assembly function only; does not impersonate the native image-generation workflow.
p=out/'ppt';im=p/'fixture-deck/origin_image';im.mkdir(parents=True,exist_ok=True);shutil.copy2(out/'launch-figure.png',im/'slide_01.png')
(p/'fixture-deck/speech.md').write_text('## Slide 1\nLocal software fixture: expired controlled launches 0 of 6; active 6 of 6. No claim of statistical generalization.\n')
r=run('codex-ppt','skills/codex-ppt/scripts/assemble_ppt.py',[p,'fixture-deck.pptx']);assert r.returncode==0,r.stderr
from pptx import Presentation
prs=Presentation(p/'fixture-deck/fixture-deck.pptx');assert len(prs.slides)==1;assert len(prs.slides[0].shapes)==1
assert '0 of 6' in prs.slides[0].notes_slide.notes_text_frame.text
(out/'ppt-inspection.json').write_text(json.dumps({'slides':1,'picture_shapes':1,'speaker_notes_verified':True,'editable_text_verified':False,'native_image_generation_executed':False},indent=2))
# Claim/evidence checking with both valid artifact and a metric-only negative control.
p=out/'spine';p.mkdir(exist_ok=True)
(p/'results_validation.md').write_text('''# Results as validation\n\n| Results Unit | Contribution Claim Tested | Result/Evidence | Allowed Interpretation | Interpretation NOT Allowed |\n|---|---|---|---|---|\n| Deadline launches | C1 expired work is denied | launch-results.csv: controlled expired 0/6 and active 6/6 | Local fixture works | Guarantees every host API call stops |\n| Source work | C2 real source workflows produce task outputs | kdense-profile.json and launch-figure.png | Selected local entry paths work | All repository capabilities work on Mac |\n''')
r=run('paperspine-workbench','dist/codex/skills/paper-spine/scripts/results_validation_check.py',[p,'--json','--write']);assert r.returncode==0
bad=p/'negative';bad.mkdir(exist_ok=True);(bad/'results_validation.md').write_text('| Results Unit | Contribution Claim Tested | Result/Evidence |\n|---|---|---|\n| Bad metric | | 0.95 |\n')
r=run('paperspine-workbench','dist/codex/skills/paper-spine/scripts/results_validation_check.py',[bad,'--json']);assert r.returncode==1
(out/'artifact-command-receipts.json').write_text(json.dumps(receipts,ensure_ascii=False,indent=2))
print(json.dumps({'rougier':True,'ppt_assembly':True,'spine_positive_and_negative':True}))
