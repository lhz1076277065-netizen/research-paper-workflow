"""Use actual evaluation rows as manuscript build data, then generate links."""
import hashlib, importlib.util, json, statistics, textwrap
from decimal import Decimal, ROUND_HALF_EVEN
from pathlib import Path
import numpy as np
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

HERE=Path(__file__).resolve().parent
SCRIPT=HERE.parents[2]/'src/common/scripts/result_links.py'
spec=importlib.util.spec_from_file_location('actual_build_links',SCRIPT);audit=importlib.util.module_from_spec(spec);spec.loader.exec_module(audit)

def main():
    independent=HERE/'independent';out=HERE/'manuscript-build';out.mkdir(exist_ok=True)
    evaluation=json.loads((independent/'evaluation.json').read_text());rng=np.random.default_rng(78534)
    results={};build=[];doc=Document();pdf=canvas.Canvas(str(out/'study.pdf'),pagesize=A4)
    pdf_lines=[];y=775;lines=['# Frozen holdout: streaming-method comparison',''];paragraph=1
    doc.add_heading('Frozen holdout: streaming-method comparison',level=1)
    pdf.setFont('Helvetica-Bold',13);pdf.drawString(35,813,'Frozen holdout: streaming-method comparison')
    pdf.setFont('Helvetica',8);pdf.drawString(35,798,'480 synthetic streams; task-level method evidence; publication novelty unestablished.')
    for name in ('A1','B1','B2','A2'):
        report=json.loads((independent/(name+'-refined.json')).read_text());loss=np.array([r['loss'] for r in report['rows']])
        ci=np.quantile(rng.choice(loss,size=(2000,len(loss))).mean(axis=1),[.025,.975])
        rid=name+'refined';base=evaluation['participants'][name]['baseline']['mean_loss']
        record={'value':report['mean_loss'],'lower':float(ci[0]),'upper':float(ci[1]),
            'unit':'loss-points','direction':'decrease' if report['mean_loss']<base else 'increase',
            'denominator':len(loss),'outcome':'penalty loss','population':'synthetic streams','time':'360 steps',
            'comparison':'own tuned baseline','model':rid,'split':'holdout','uncertainty':'bootstrap interval',
            'effect_type':'arithmetic mean'}
        results[rid]=record
        def display(field):return format(Decimal(str(record[field])).quantize(Decimal('.01'),rounding=ROUND_HALF_EVEN),'.2f')
        text=(f'{rid} = {display("value")} loss-points, bootstrap interval '
              f'[{display("lower")} loss-points, {display("upper")} loss-points]. '
              f'n={len(loss)}; penalty loss; synthetic streams; 360 steps; own tuned baseline; '
              f'holdout; arithmetic mean; direction={record["direction"]}.')
        short=f'{rid} = {display("value")} loss-points.'
        for role,body,fields in [('body',text,('value','lower','upper')),('caption',short,('value',))]:
            numeric=[{'field':field,'decimals':2} for field in fields]
            lines.append(body);build.append({'path':'study.md','locator':{'line':len(lines)},'result_id':rid,'role':role,'numeric':numeric})
            doc.add_paragraph(body);paragraph+=1
            build.append({'path':'study.docx','locator':{'paragraph':paragraph},'result_id':rid,'role':role,'numeric':numeric})
            chunks=textwrap.wrap(body,width=88,break_long_words=False,break_on_hyphens=False)
            start=len(pdf_lines)+1
            for chunk in chunks:
                pdf.setFont('Helvetica',9);pdf.drawString(35,y,chunk);y-=13;pdf_lines.append(chunk)
            build.append({'path':'study.pdf','locator':{'page':1,'line_start':start,'line_end':len(pdf_lines)},'result_id':rid,'role':role,'numeric':numeric})
            y-=9
    pdf.setFont('Helvetica',8);pdf.drawString(35,28,'Paired synthetic cases; bootstrap intervals are exploratory. Page 1')
    pdf.save();doc.save(out/'study.docx');(out/'study.md').write_text('\n'.join(lines)+'\n')
    (out/'results.json').write_text(json.dumps(results,indent=2)+'\n')
    (out/'execution.json').write_text(json.dumps({'evaluation_sha256':hashlib.sha256((independent/'evaluation.json').read_bytes()).hexdigest(),
        'input_holdout_sha256':evaluation['holdout_sha256'],'unseen_evaluation':True,'renderer':'Actual post-freeze row reports -> Markdown/DOCX/PDF'},indent=2)+'\n')
    # Existing build rows supply actual paths, roles, locations and display precision.
    # PDF coordinates come from the actual extractor output, not canvas row counts.
    page_lines=audit._pdf_pages(out/'study.pdf')[0].splitlines()
    pdf_rows=[row for row in build if row['path']=='study.pdf'];starts=[];cursor=0
    for row in pdf_rows:
        first=next(i for i in range(cursor,len(page_lines)) if row['result_id']+' =' in page_lines[i])
        starts.append(first);cursor=first+1
    for i,row in enumerate(pdf_rows):
        end=starts[i+1] if i+1<len(starts) else len(page_lines)
        row['locator']={'page':1,'line_start':starts[i]+1,'line_end':end}
    sources=[{'id':'study','path':'results.json','versions':{'data':{'path':'../independent/holdout.json'},
        'code':{'path':'../independent_eval.py'},'execution':{'path':'execution.json'}}}]
    # root is the full study project, so all file paths remain project-relative.
    for source in sources:
        source['path']='manuscript-build/'+source['path']
        source['versions']={'data':{'path':'independent/holdout.json'},'code':{'path':'independent_eval.py'},'execution':{'path':'manuscript-build/execution.json'},'method_snapshots':{'path':'independent/frozen-submissions.json'},'render':{'path':'render_actual.py'}}
    for row in build:row['path']='manuscript-build/'+row['path']
    (out/'build-rows.json').write_text(json.dumps(build,indent=2)+'\n')
    payload=audit.build_links(HERE,sources,build);report=audit.audit_links(payload,HERE)
    (out/'links.json').write_text(json.dumps(payload,indent=2)+'\n');(out/'result-links-report.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps({'passed':report['passed'],'occurrences':len(report['coverage']['occurrences']),'pending':report['pending'],'errors':report['errors']}))
    if not report['passed']:raise AssertionError(report)

if __name__=='__main__':main()
