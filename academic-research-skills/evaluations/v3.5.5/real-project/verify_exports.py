from pathlib import Path
import json,zipfile,re,hashlib
from docx import Document
from pypdf import PdfReader
P=Path(__file__).resolve().parent;doc=Document(P/'wine-association-paper.docx');pdf=PdfReader(P/'wine-association-paper.pdf')
md=(P/'manuscript-final.md').read_text();body=[p.text for p in doc.paragraphs]
expected=[l for l in md.splitlines() if l.strip() and not l.startswith(('#','!','|'))]
assert all(l in body for l in expected)
assert len(doc.tables)==1 and len(doc.tables[0].rows)==6
for line,row in zip([l for l in md.splitlines() if l.startswith('|') and not re.match(r'^\|[-| ]+\|$',l)],doc.tables[0].rows):assert [x.strip() for x in line.strip('|').split('|')]==[c.text for c in row.cells]
with zipfile.ZipFile(P/'wine-association-paper.docx') as z:
 assert z.testzip() is None;media=[x for x in z.namelist() if x.startswith('word/media/')];assert len(media)==3
 assert b'pBdr' not in z.read('word/document.xml')
assert len(pdf.pages)==5
alltext='\n'.join(x.extract_text() for x in pdf.pages)
for required in ['0.546','1.560','0.940','0.319','1.094','0.464','1.527','1.246','1.875','0.492','1.389','178','59','71','48','Methods','Results','Discussion','References','Figure 1.','Figure 2.','Supplementary Figure S1.']:
 assert required in alltext,required
assert '[3]' not in alltext and 'installing an additional statistical stack' not in alltext
for n in range(1,6):assert (P/'rendered'/f'page-{n}.png').stat().st_size>10000
r={'docx_opened_and_ooxml_integrity':True,'all_manuscript_prose_paragraphs_exact_in_docx':True,'editable_table_rows':6,'embedded_figures':len(media),'pdf_real_renderer':'Canonical DOCX renderer using bundled LibreOffice and Poppler','pdf_pages':len(pdf.pages),'pdf_key_scientific_numbers_and_sections_present':True,'actual_visual_review':'All5final PNG pages opened and inspected after final content rebuild. No clipped/overlapping body text, broken table rows, missing glyphs or figure/interval clipping. Arial title black with inherited border removed; captions normal weight; all figures remain exact6.5inchwidth. Page breaks do not remove any paragraph.','visual_revision_history':'Figure1initialfootnoteoverlap fixed; inherited Wordtitleblueborder removed; before-editorial and failed duplicate-referencecheck retained. Finaldocx/PDF rerendered after allchanges.','scientific_changes':'No numericalresult/data/modelchanges. All actualv2->final textual changes assessed in expression-meaning-checks.json. Proceduralmetadata/thirdreference moved to README.','known_limits':['No accountablehuman authorapproval or externalpeerreview.','Source stat/profile/citationrequests helpers lacking optional dependencies were not claimed successful; explicitadaptations/sourcefunctions documented.','Five-page report includes3textpages,1mainfigurepage,1supplementfigurepage.'],'artifacts':{f:{'bytes':(P/f).stat().st_size,'sha256':hashlib.sha256((P/f).read_bytes()).hexdigest()} for f in ['manuscript-final.md','wine-association-paper.docx','wine-association-paper.pdf']}}
(P/'export-verification-final.json').write_text(json.dumps(r,indent=2))
print(json.dumps(r,indent=2))
