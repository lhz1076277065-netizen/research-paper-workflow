from pathlib import Path
import re,json,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parent;D=Document();sec=D.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.35)
for n in ['Normal','Body Text','Title','Heading 1','Heading 2','Caption']:
 st=D.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.10
D.styles['Normal'].font.size=Pt(10.5);D.styles['Title'].font.size=Pt(17);D.styles['Title'].font.bold=True
D.styles['Title'].paragraph_format.space_after=Pt(12)
D.styles['Heading 1'].font.size=Pt(12);D.styles['Heading 1'].paragraph_format.space_before=Pt(10);D.styles['Heading 1'].paragraph_format.space_after=Pt(5)
D.styles['Heading 1'].paragraph_format.keep_with_next=True
D.styles['Caption'].font.size=Pt(9);D.styles['Caption'].paragraph_format.line_spacing=1.05
h=sec.header.paragraphs[0];h.text='UCI Wine descriptive association';h.style='Normal';h.runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run();r.font.size=Pt(8);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
lines=(P/'manuscript-final.md').read_text().splitlines();i=0
while i<len(lines):
 t=lines[i];i+=1
 if not t.strip():continue
 if t.startswith('# '):D.add_paragraph(t[2:],'Title');continue
 if t.startswith('## '):
  title=t[3:]
  if title=='Figures':D.add_page_break()
  if title=='Supplementary results':D.add_page_break()
  D.add_paragraph(title,'Heading 1');continue
 if t.startswith('|'):
  rows=[t]
  while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
  rows=[x for x in rows if not re.match(r'^\|[-| ]+\|$',x)]
  values=[[c.strip() for c in row.strip('|').split('|')] for row in rows]
  table=D.add_table(rows=len(values),cols=len(values[0]));table.autofit=False
  widths=[1.5,.45,.95,1,.75,1.85]
  for rr,v in enumerate(values):
   for cc,txt in enumerate(v):
    cell=table.cell(rr,cc);cell.width=Inches(widths[cc]);cell.text=txt;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     node=OxmlElement('w:'+edge);node.set(qn('w:w'),'70');node.set(qn('w:type'),'dxa');mar.append(node)
    tcpr.append(mar);borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
     node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
    tcpr.append(borders)
    if rr==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');tcpr.append(sh)
    for p in cell.paragraphs:
     p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
     if cc:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
     for r in p.runs:r.font.size=Pt(9);r.font.bold=(rr==0)
   pr=table.rows[rr]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
  hd=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(hd)
  D.add_paragraph();continue
 if t.startswith('!['):
  match=re.match(r'!\[([^]]+)\]\(([^)]+)\)',t);p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  run=p.add_run();run.add_picture(str(P/match.group(2)),width=Inches(6.5))
  docpr=run._r.xpath('.//wp:docPr')[0];docpr.set('descr',match.group(1)+' from complete official UCI data; editable SVG counterpart included')
  continue
 if t.startswith(('Figure ','Supplementary Figure ','Table ')):
  p=D.add_paragraph(t,'Caption');p.paragraph_format.keep_with_next=t.startswith('Table ');continue
 p=D.add_paragraph(t)
 if t.startswith('['):
  for r in p.runs:r.font.size=Pt(9)
D.core_properties.title='Alcohol content and color intensity in the UCI Wine dataset';D.core_properties.subject='Complete historical dataset descriptive association';D.core_properties.author='';D.core_properties.comments='AI-assisted research report; no author approval or external peer review asserted.'
D.save(P/'wine-association-paper.docx')
with zipfile.ZipFile(P/'wine-association-paper.docx') as z:assert z.testzip() is None
(P/'export-build.json').write_text(json.dumps({'docx':'wine-association-paper.docx','editable_text':True,'editable_table':True,'figure_count':3,'image_width_inches':6.5,'figure_raster_dpi':300,'vector_sources':'figures/*.svg and *.pdf','source':'manuscript-final.md','python_docx':'1.2.0','no_invented_author':True},indent=2))
print('DOCX built')
